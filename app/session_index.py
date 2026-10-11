"""Load uploaded files (PDF, DOCX, PPTX), tag them, chunk them and build a FAISS index."""
import os
import tempfile

from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

SOURCE_TYPES = ["notes", "textbook", "question_paper", "lab_manual"]
SUPPORTED_EXTENSIONS = ["pdf", "docx", "pptx"]


def _load_pptx(path: str) -> list[Document]:
    from pptx import Presentation

    docs = []
    for number, slide in enumerate(Presentation(path).slides, start=1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if texts:
            docs.append(Document(page_content="\n".join(texts), metadata={"slide": number}))
    return docs


def load_file(file_name: str, file_bytes: bytes, source_type: str, subject: str = "", chapter: str = "") -> list[Document]:
    """Read one uploaded file and tag every piece with its file name, type, subject and chapter."""
    extension = file_name.rsplit(".", 1)[-1].lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: .{extension}")

    # Loaders need a real file on disk, so write the upload to a temp file first.
    with tempfile.NamedTemporaryFile(delete=False, suffix=f".{extension}") as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
    try:
        if extension == "pdf":
            docs = PyPDFLoader(tmp_path).load()
        elif extension == "docx":
            docs = Docx2txtLoader(tmp_path).load()
        else:
            docs = _load_pptx(tmp_path)
    finally:
        os.remove(tmp_path)

    for doc in docs:
        doc.metadata["source"] = file_name
        doc.metadata["source_type"] = source_type
        if subject.strip():
            doc.metadata["subject"] = subject.strip()
        if chapter.strip():
            doc.metadata["chapter"] = chapter.strip()
    return docs


def build_session_index(files, embeddings):
    """files = list of (file_name, file_bytes, source_type) or
    (file_name, file_bytes, source_type, subject, chapter). Returns (FAISS index, chunk count)."""
    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=100)
    chunks = []
    for file_name, file_bytes, source_type, *tags in files:
        subject = tags[0] if len(tags) > 0 else ""
        chapter = tags[1] if len(tags) > 1 else ""
        docs = load_file(file_name, file_bytes, source_type, subject, chapter)
        chunks.extend(splitter.split_documents(docs))
    if not chunks:
        raise ValueError("No readable text found in the uploaded files.")
    return FAISS.from_documents(chunks, embeddings), len(chunks)
