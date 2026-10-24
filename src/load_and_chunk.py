import os
import re
from langchain_community.document_loaders import PyPDFLoader, Docx2txtLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation


def load_pptx(file_path):
    """Read a PowerPoint file: one Document per slide, with the slide number."""
    prs = Presentation(file_path)
    docs = []
    for number, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        slide_text = "\n".join(t for t in texts if t.strip())
        if slide_text.strip():
            docs.append(
                Document(
                    page_content=slide_text,
                    metadata={"source": file_path, "slide": number},
                )
            )
    return docs


def read_plain_text(file_path):
    """Read a whole PDF or DOCX as one block of text."""
    lower = file_path.lower()
    if lower.endswith(".pdf"):
        pages = PyPDFLoader(file_path).load()
    elif lower.endswith(".docx"):
        pages = Docx2txtLoader(file_path).load()
    else:
        raise ValueError(f"Question papers must be PDF or DOCX: {file_path}")
    return "\n".join(p.page_content for p in pages)


def split_questions(file_path, year=None):
    """Split a question paper so every question (Q1., Q2., ...) is its own Document."""
    text = read_plain_text(file_path)
    parts = re.split(r"(?m)^\s*Q(\d+)[.)]\s*", text)
    # parts = [text before Q1, "1", text of Q1, "2", text of Q2, ...]
    docs = []
    for i in range(1, len(parts) - 1, 2):
        number = int(parts[i])
        body = parts[i + 1].strip()
        if not body:
            continue
        metadata = {"source": file_path, "question_no": number}
        if year is not None:
            metadata["year"] = int(year)
        marks_match = re.search(r"\[(\d+)\s*marks?\]", body, re.IGNORECASE)
        if marks_match:
            metadata["marks"] = int(marks_match.group(1))
        docs.append(Document(page_content=f"Q{number}. {body}", metadata=metadata))
    return docs


def load_single_file(file_path, source_type="notes", subject="AIML",
                     chapter="General", year=None):
    """Load one file and tag it with metadata."""
    lower = file_path.lower()
    if source_type == "question_paper":
        documents = split_questions(file_path, year)
    elif lower.endswith(".pdf"):
        documents = PyPDFLoader(file_path).load()
    elif lower.endswith(".docx"):
        documents = Docx2txtLoader(file_path).load()
    elif lower.endswith(".pptx"):
        documents = load_pptx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_path}")

    for doc in documents:
        doc.metadata["source_type"] = source_type
        doc.metadata["source_file"] = os.path.basename(file_path)
        doc.metadata["subject"] = subject
        doc.metadata["chapter"] = chapter

    return documents


def load_and_chunk_multiple(file_configs):
    """
    file_configs: list of dicts like
      {"path": "data/unit1.pptx", "source_type": "notes",
       "subject": "AIML", "chapter": "Unit 1 - Introduction to AI"}
    Question papers also take a "year". Each question stays as one chunk.
    """
    note_docs = []
    question_docs = []
    for config in file_configs:
        docs = load_single_file(
            config["path"],
            config["source_type"],
            config.get("subject", "AIML"),
            config.get("chapter", "General"),
            config.get("year"),
        )
        if config["source_type"] == "question_paper":
            question_docs.extend(docs)
        else:
            note_docs.extend(docs)

    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150)
    chunks = splitter.split_documents(note_docs) if note_docs else []
    return chunks + question_docs


# Keep the old single-file function for backward compatibility
def load_and_chunk(pdf_path):
    return load_and_chunk_multiple([{"path": pdf_path, "source_type": "notes"}])


if __name__ == "__main__":
    questions = split_questions("data/aiml_sample_qp_2023.docx", 2023)
    print(len(questions), "questions found")
    print(questions[0].metadata)
    print(questions[0].page_content)